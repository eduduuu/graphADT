#!/usr/bin/env python3
"""
Script para processamento de resultados e geração de Dashboards (Excel, HTMLs e CSV).
Atualizado para o formato unificado de micro_benchmark.csv e exibição de Total Throughput.
"""

import os
import sys
import numpy as np
import pandas as pd
import scipy.stats as stats
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
import plotly.graph_objects as go

def process_benchmark_data(input_csv="micro_benchmark.csv", output_ci_csv="confidence_interval.csv"):
    if not os.path.exists(input_csv):
        raise FileNotFoundError(f"Arquivo de entrada não encontrado: {input_csv}")

    df = pd.read_csv(input_csv)
    print(f"[1/4] Carregado {input_csv} ({len(df)} execuções)")

    # Normalizar nomes de colunas para maiúsculas compatíveis
    col_rename = {}
    if 'implementation' in df.columns:
        col_rename['implementation'] = 'Implementation'
    if 'workload' in df.columns:
        col_rename['workload'] = 'Workload'
    df.rename(columns=col_rename, inplace=True)

    # Identificar coluna de buscas
    search_col = 'e_searches_done' if 'e_searches_done' in df.columns else 'e_searches'
    v_ins = df['v_insert'] if 'v_insert' in df.columns else 0
    v_del = df['v_del'] if 'v_del' in df.columns else 0

    # Total de operações e Throughput Total (ops/s)
    total_ops = v_ins + v_del + df['e_insert'] + df['e_del'] + df[search_col]
    duration_sec = df['duration'] / 1000.0
    df['throughput_total'] = total_ops / duration_sec

    # Agrupar e calcular Média e Intervalo de Confiança (95%)
    group_cols = ['size', 'duration', 'density', 'Implementation', 'Workload']
    summary_data = []

    for group_name, group_df in df.groupby(group_cols):
        group_stats = dict(zip(group_cols, group_name))
        n = len(group_df)
        mean_val = group_df['throughput_total'].mean()

        if n < 2:
            ci_l, ci_u = mean_val, mean_val
        else:
            se = stats.sem(group_df['throughput_total'])
            ci_l, ci_u = stats.t.interval(0.95, df=n-1, loc=mean_val, scale=se)
            if np.isnan(ci_l): ci_l = mean_val
            if np.isnan(ci_u): ci_u = mean_val

        group_stats['throughput_total_mean'] = mean_val
        group_stats['throughput_total_ci_lower'] = max(0.0, float(ci_l))
        group_stats['throughput_total_ci_upper'] = float(ci_u)
        group_stats['n_samples'] = n
        summary_data.append(group_stats)

    summary_df = pd.DataFrame(summary_data)
    summary_df.to_csv(output_ci_csv, index=False)
    print(f"[2/4] Salvo {output_ci_csv} ({len(summary_df)} configurações agrupadas)")
    return summary_df


def build_excel_dashboard(summary_df, output_excel="Experiment_Dashboard.xlsx"):
    workload_order = ['read_only', 'mostly_read', 'mostly_write']

    def get_sorted_vals(col):
        vals = summary_df[col].unique()
        if col == 'Workload':
            return [w for w in workload_order if w in vals]
        return sorted(vals)

    def build_sheet(wb, sheet_name, row_col, group_col, x_col, temp_images):
        ws = wb.create_sheet(title=sheet_name)
        ws.sheet_view.showGridLines = False

        # Título
        ws.cell(row=1, column=2, value=f"{sheet_name} Dashboard").font = Font(size=20, bold=True, color="1F4E79")
        ws.cell(row=2, column=2, value=f"X: {x_col.capitalize()} | Y: Mean Throughput (log ops/sec) | CI 95%").font = Font(size=11, italic=True)

        current_row = 4
        row_vals = get_sorted_vals(row_col)
        group_vals = get_sorted_vals(group_col)

        for r_val in row_vals:
            ws.cell(row=current_row, column=2, value=f"{row_col.capitalize()}: {r_val}").font = Font(size=14, bold=True, color="FFFFFF")
            ws.cell(row=current_row, column=2).fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            current_row += 2

            for g_val in group_vals:
                sub_df = summary_df[(summary_df[row_col] == r_val) & (summary_df[group_col] == g_val)].copy()
                if sub_df.empty:
                    continue

                ws.cell(row=current_row, column=2, value=f"{group_col.capitalize()}: {g_val}").font = Font(size=12, bold=True, color="1F4E79")
                ws.cell(row=current_row, column=2).fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
                current_row += 2

                fig, ax = plt.subplots(figsize=(6.5, 4.2))
                for impl in sub_df['Implementation'].unique():
                    impl_df = sub_df[sub_df['Implementation'] == impl].copy()
                    if x_col == 'Workload':
                        impl_df['Workload'] = pd.Categorical(impl_df['Workload'], categories=workload_order, ordered=True)
                    impl_df = impl_df.sort_values(x_col)

                    x = impl_df[x_col].astype(str)
                    y = impl_df['throughput_total_mean']
                    ci_l = impl_df['throughput_total_ci_lower']
                    ci_u = impl_df['throughput_total_ci_upper']

                    yerr = [np.maximum(0, y - ci_l), np.maximum(0, ci_u - y)]
                    ax.errorbar(x, y, yerr=yerr, fmt='-o', label=impl, capsize=4, capthick=1.2, markersize=5)

                ax.set_title("Total Throughput", fontsize=11, fontweight='bold')
                ax.set_xlabel(x_col.capitalize(), fontsize=10)
                ax.set_ylabel('Throughput (ops/s)', fontsize=10)
                ax.set_yscale('log')
                ax.legend(title="Implementation", fontsize=9)
                ax.grid(True, linestyle='--', alpha=0.5)
                plt.tight_layout()

                safe_r = str(r_val).replace('.', '_').replace('/', '_')
                safe_g = str(g_val).replace('.', '_').replace('/', '_')
                temp_img_path = f"temp_{sheet_name.replace(' ', '')}_{safe_r}_{safe_g}.png"
                fig.savefig(temp_img_path, dpi=100)
                plt.close(fig)
                temp_images.append(temp_img_path)

                ws.add_image(Image(temp_img_path), f"B{current_row}")
                current_row += 23
            current_row += 1

    wb = Workbook()
    wb.remove(wb.active) # Remove sheet padrão
    temp_images = []

    build_sheet(wb, "By Size and Density", row_col='size', group_col='density', x_col='Workload', temp_images=temp_images)
    build_sheet(wb, "By Workload and Density", row_col='Workload', group_col='density', x_col='size', temp_images=temp_images)
    build_sheet(wb, "By Workload and Size", row_col='Workload', group_col='size', x_col='density', temp_images=temp_images)

    wb.save(output_excel)
    for img in temp_images:
        if os.path.exists(img):
            os.remove(img)
    print(f"[3/4] Salvo {output_excel}")


def build_html_dashboards(df, output_prefix="dashboard"):
    workload_mapping = {'read_only': 1, 'mostly_read': 2, 'mostly_write': 3}
    rev_workload_mapping = {1: 'Read Only', 2: 'Mostly Read', 3: 'Mostly Write'}
    df = df.copy()
    df['Workload_num'] = df['Workload'].map(workload_mapping)

    implementations = ['CSR', 'ADJACENCY_MATRIX', 'ADJACENCY_LIST']
    impl_colors = {
        'CSR': 'rgba(31, 119, 180, 0.7)',
        'ADJACENCY_MATRIX': 'rgba(214, 39, 40, 0.7)',
        'ADJACENCY_LIST': 'rgba(44, 160, 44, 0.7)'
    }

    def build_single_3d_plot(filtered_df, x_col, y_col, x_label, y_label):
        fig = go.Figure()

        for impl in implementations:
            impl_df = filtered_df[filtered_df['Implementation'] == impl]
            if impl_df.empty:
                continue

            z_values = impl_df['throughput_total_mean'].clip(lower=1)

            if len(impl_df) >= 3:
                fig.add_trace(go.Mesh3d(
                    x=impl_df[x_col], y=impl_df[y_col], z=z_values,
                    opacity=0.55, color=impl_colors[impl], name=f"{impl} (Superfície)",
                    showlegend=True
                ))

            fig.add_trace(go.Scatter3d(
                x=impl_df[x_col], y=impl_df[y_col], z=z_values,
                mode='markers',
                marker=dict(size=5, color=impl_colors[impl].replace('0.7', '1.0')),
                name=impl,
                text=impl_df['throughput_total_mean'],
                hovertemplate=(
                    f"<b>{impl}</b><br>"
                    f"{x_label}: %{{x}}<br>"
                    f"{y_label}: %{{y}}<br>"
                    "Throughput: %{text:,.1f} ops/s<extra></extra>"
                )
            ))

        xaxis_dict = dict(title=x_label)
        yaxis_dict = dict(title=y_label)
        if x_col == 'Workload_num':
            xaxis_dict.update(tickvals=[1, 2, 3], ticktext=['RO', 'MR', 'MW'])
        if y_col == 'Workload_num':
            yaxis_dict.update(tickvals=[1, 2, 3], ticktext=['RO', 'MR', 'MW'])

        fig.update_layout(
            scene=dict(
                xaxis=xaxis_dict,
                yaxis=yaxis_dict,
                zaxis=dict(title='Total Throughput (Log Scale)', type='log')
            ),
            height=580,
            margin=dict(l=10, r=10, b=10, t=30),
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
        )
        return fig.to_html(full_html=False, include_plotlyjs=False)

    def get_html_header(title):
        return f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{title}</title>
    <script src="https://cdn.plot.ly/plotly-2.32.0.min.js"></script>
    <style>
        body {{ font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; background: #f8f9fa; margin: 0; padding: 20px; }}
        .header-box {{ background: #0d436b; color: white; padding: 15px; border-radius: 8px; margin: 30px 0 15px 0; text-align: center; }}
        .chart-container {{ background: white; border-radius: 8px; box-shadow: 0 4px 12px rgba(0,0,0,0.1); padding: 10px; margin-bottom: 30px; }}
    </style>
</head>
<body>
    <h1 style="text-align: center; color: #222;">{title}</h1>
"""

    # Dashboard 1: Fixado por Tamanho
    html_1 = get_html_header("Dashboard 1: Fixado por Tamanho (X=Density, Y=Workload)")
    for sz in sorted(df['size'].unique()):
        html_1 += f'<div class="header-box"><h2>Size = {sz}</h2></div><div class="chart-container">'
        html_1 += build_single_3d_plot(df[df['size'] == sz], 'density', 'Workload_num', 'Density', 'Workload')
        html_1 += '</div>'
    html_1 += "</body></html>"
    with open("dashboard_1_size.html", "w", encoding="utf-8") as f:
        f.write(html_1)

    # Dashboard 2: Fixado por Densidade
    html_2 = get_html_header("Dashboard 2: Fixado por Densidade (X=Workload, Y=Size)")
    for d in sorted(df['density'].unique()):
        html_2 += f'<div class="header-box"><h2>Density = {d}</h2></div><div class="chart-container">'
        html_2 += build_single_3d_plot(df[df['density'] == d], 'Workload_num', 'size', 'Workload', 'Size')
        html_2 += '</div>'
    html_2 += "</body></html>"
    with open("dashboard_2_density.html", "w", encoding="utf-8") as f:
        f.write(html_2)

    # Dashboard 3: Fixado por Carga de Trabalho
    html_3 = get_html_header("Dashboard 3: Fixado por Carga de Trabalho (X=Size, Y=Density)")
    for w in sorted(df['Workload'].unique()):
        w_title = rev_workload_mapping[workload_mapping[w]]
        html_3 += f'<div class="header-box"><h2>Workload = {w_title}</h2></div><div class="chart-container">'
        html_3 += build_single_3d_plot(df[df['Workload'] == w], 'size', 'density', 'Size', 'Density')
        html_3 += '</div>'
    html_3 += "</body></html>"
    with open("dashboard_3_workload.html", "w", encoding="utf-8") as f:
        f.write(html_3)

    # Dashboard All-in-One (com Jitter nos pontos e escala de cores por Throughput Total)
    impl_config = {
        'CSR': {'symbol': 'circle', 'offset': -0.15},
        'ADJACENCY_MATRIX': {'symbol': 'square', 'offset': 0.0},
        'ADJACENCY_LIST': {'symbol': 'diamond', 'offset': 0.15}
    }
    valid_data = df['throughput_total_mean'].clip(lower=1)
    cmin = np.log10(valid_data.min())
    cmax = np.log10(valid_data.max())

    fig_all = go.Figure()
    for impl, config in impl_config.items():
        plot_df = df[df['Implementation'] == impl].copy()
        if not plot_df.empty:
            log_tp = np.log10(plot_df['throughput_total_mean'].clip(lower=1))
            jittered_x = plot_df['Workload_num'] + config['offset']
            fig_all.add_trace(go.Scatter3d(
                x=jittered_x, y=plot_df['size'], z=plot_df['density'],
                mode='markers',
                marker=dict(
                    symbol=config['symbol'],
                    size=9,
                    color=log_tp,
                    colorscale='RdYlGn',
                    cmin=cmin,
                    cmax=cmax,
                    showscale=(impl == 'ADJACENCY_LIST'),
                    colorbar=dict(title="Throughput<br>(Log Scale)", x=1.02) if (impl == 'ADJACENCY_LIST') else None,
                    line=dict(width=1, color='black'),
                    opacity=1.0
                ),
                name=impl,
                text=plot_df['throughput_total_mean'],
                hovertemplate=(
                    f"<b>{impl}</b><br>"
                    "Workload: %{x}<br>"
                    "Size: %{y}<br>"
                    "Density: %{z}<br>"
                    "Throughput: %{text:,.1f} ops/s<extra></extra>"
                )
            ))

    fig_all.update_layout(
        scene=dict(
            xaxis=dict(title='Workload', tickvals=[1, 2, 3], ticktext=['RO', 'MR', 'MW'], range=[0.5, 3.5]),
            yaxis=dict(title='Size', type='category'),
            zaxis=dict(title='Density'),
            camera=dict(eye=dict(x=1.6, y=1.6, z=1.2))
        ),
        height=680,
        margin=dict(l=10, r=10, b=10, t=40),
        legend=dict(orientation="h", yanchor="bottom", y=1.05, xanchor="right", x=1)
    )

    html_all = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>Global Architecture Comparison (Total Throughput)</title>
    <script src="https://cdn.plot.ly/plotly-2.32.0.min.js"></script>
    <style>
        body {{ font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; background: #f8f9fa; margin: 0; padding: 20px; }}
        .header-box {{ background: #0d436b; color: white; padding: 15px; border-radius: 8px; margin: 10px 0 20px 0; text-align: center; }}
        .legend-bar {{ text-align: center; font-weight: bold; margin-bottom: 20px; padding: 15px; background: #e9ecef; border-radius: 8px; font-size: 15px; line-height: 1.6; }}
        .chart-container {{ background: white; border-radius: 8px; box-shadow: 0 4px 12px rgba(0,0,0,0.1); padding: 10px; }}
    </style>
</head>
<body>
    <div class="header-box">
        <h1 style="margin:0;">Global Architecture Comparison (Total Throughput)</h1>
    </div>
    <div class="legend-bar">
        🟢 <b>Verde</b> = Alto Throughput (Rápido) &nbsp;|&nbsp; 🟡 <b>Amarelo</b> = Médio &nbsp;|&nbsp; 🔴 <b>Vermelho</b> = Baixo Throughput (Lento)<br>
        ● <b>Círculo:</b> CSR &nbsp;|&nbsp; ■ <b>Quadrado:</b> Matriz de Adjacência &nbsp;|&nbsp; ◆ <b>Losango:</b> Lista de Adjacência<br>
        <span style="font-size:13px; font-weight:normal; color:#555;">(Símbolos deslocados no eixo Workload para evitar sobreposição)</span>
    </div>
    <div class="chart-container">
        {fig_all.to_html(full_html=False, include_plotlyjs=False)}
    </div>
</body>
</html>
"""
    with open("dashboard_all_in_one.html", "w", encoding="utf-8") as f:
        f.write(html_all)

    print("[4/4] Salvo dashboard_1_size.html, dashboard_2_density.html, dashboard_3_workload.html e dashboard_all_in_one.html")


if __name__ == "__main__":
    input_file = sys.argv[1] if len(sys.argv) > 1 else "micro_benchmark.csv"
    summary_df = process_benchmark_data(input_file)
    build_excel_dashboard(summary_df)
    build_html_dashboards(summary_df)
    print("\nTodos os dashboards foram gerados com sucesso!")
